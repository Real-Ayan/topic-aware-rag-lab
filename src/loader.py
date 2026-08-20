"""Load PDF, DOCX, and plain-text files from data/raw/ (recursive)."""

from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from langchain_community.document_loaders import PyMuPDFLoader
from rich.console import Console

from src.config import settings

console = Console()

SUPPORTED_SUFFIXES = {".pdf", ".txt", ".md", ".docx"}
SKIP_NAMES = {"readme.md", ".gitkeep"}


def _rel_source(path: Path, root: Path) -> str:
    """Stable source label relative to data root (keeps nested folder context)."""
    try:
        return str(path.relative_to(root)).replace("\\", "/")
    except ValueError:
        return path.name


def _discover_files(directory: Path) -> list[Path]:
    """Recursively find all supported files under directory."""
    files: list[Path] = []
    for path in sorted(directory.rglob("*")):
        if not path.is_file():
            continue
        if path.name.lower() in SKIP_NAMES:
            continue
        if path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        # Skip hidden / system junk
        if any(part.startswith(".") for part in path.relative_to(directory).parts):
            continue
        files.append(path)
    return files


def _load_pdf(path: Path, source: str) -> list[dict]:
    docs = PyMuPDFLoader(str(path)).load()
    records: list[dict] = []
    for doc in docs:
        page = int(doc.metadata.get("page", 0))
        text = (doc.page_content or "").strip()
        if not text:
            continue
        records.append({"text": text, "source": source, "page": page})
    return records


def _load_text(path: Path, source: str) -> list[dict]:
    text = path.read_text(encoding="utf-8", errors="replace").strip()
    if not text:
        return []
    return [{"text": text, "source": source, "page": 0}]


def _load_docx(path: Path, source: str) -> list[dict]:
    document = DocxDocument(str(path))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))
    text = "\n\n".join(paragraphs).strip()
    if not text:
        return []
    return [{"text": text, "source": source, "page": 0}]


def load_raw_documents(raw_dir: str | Path | None = None) -> list[dict]:
    """Return flat records: {text, source, page} from all nested files."""
    directory = Path(raw_dir or settings.DATA_RAW_DIR)
    if not directory.exists():
        raise FileNotFoundError(f"Raw data directory not found: {directory}")

    files = _discover_files(directory)
    if not files:
        raise FileNotFoundError(
            f"No supported files under {directory} (searched recursively). "
            f"Supported: {', '.join(sorted(SUPPORTED_SUFFIXES))}"
        )

    console.print(
        f"[cyan]Found {len(files)} file(s) under[/cyan] {directory} "
        f"(recursive search)"
    )

    records: list[dict] = []
    empty = 0
    for path in files:
        source = _rel_source(path, directory)
        suffix = path.suffix.lower()
        try:
            if suffix == ".pdf":
                batch = _load_pdf(path, source)
            elif suffix == ".docx":
                batch = _load_docx(path, source)
            else:
                batch = _load_text(path, source)
        except Exception as exc:
            console.print(f"[red]Failed[/red] {source}: {exc}")
            continue

        if not batch:
            empty += 1
            console.print(f"  [yellow]empty[/yellow] {source}")
            continue

        console.print(f"  loaded {source}: {len(batch)} page/section(s)")
        records.extend(batch)

    console.print(
        f"[green]Total records:[/green] {len(records)} from {len(files)} file(s)"
        + (f" ({empty} empty skipped)" if empty else "")
    )
    if not records:
        raise RuntimeError(f"Files were found under {directory}, but none yielded text.")
    return records
